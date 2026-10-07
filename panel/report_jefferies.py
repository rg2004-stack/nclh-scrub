"""A Jefferies / Tourism Economics-compatible workbook, built beside the panel's own.

    python -m panel.report_jefferies --tier weekly-full

Writes `reports/<scrape_date>__<tier>__jefferies-basis.xlsx`. Nothing in the
existing workbook, its sheets or its calculations is changed; this module only
reads the panel and reuses the report's loaders, filters and table writer.

What the sell-side series is, as stated in the reports
------------------------------------------------------
"Minimum advertised price in dollars for a balcony cabin", by brand x region x
SAIL MONTH, with columns for the PRICE MONTH in which prices were observed.
Each value is the minimum over every sailing departing in that sail month and
every observation taken in that price month. No length filter, no matched
basket: whatever was listed.

The panel is built the other way round -- a mean over a fixed basket of cabins
-- precisely because a minimum confounds price with inventory: when the
cheapest cabins sell, the minimum rises and reads as pricing strength. This
workbook reproduces the minimum so the two can be put side by side, and keeps
the matched-basket minimum beside it so the difference stays measurable.

One fact about the data shapes everything below
-----------------------------------------------
After the peer-comparable exclusion (NCL's MINISUITE is filed under `balcony`
but is not a balcony), every sailing carries exactly ONE priced balcony cell on
both lines -- Carnival's `OB`, NCL's `BALCONY`. So "min over cells in a sail
month" is min over SAILINGS in that sail month, which is what the sell-side
minimum is anyway. It also means a per-sailing mean/min ratio is 1.0 by
construction, so the min-vs-mean sheet compares across sailings within a sail
month instead.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import json
import os
import sqlite3
import statistics
import sys
from datetime import date, datetime, timezone
from typing import Any, Sequence

from panel import analysis as an
from panel import report as rp
from panel.sources.capabilities import peer_excluded_subcategories

# -- region mapping ---------------------------------------------------------

# Our region -> the sell-side region. Our Caribbean already includes Bahamas on
# both lines (NCL maps its BAHAMAS destination and Carnival its BH region code
# to Caribbean), so it maps to the combined Caribbean/Bahamas one to one.
JEFFERIES_REGION: dict[str, str | None] = {
    "Caribbean": "Caribbean/Bahamas",
    "Bermuda": "Bermuda",
    "Alaska": "Alaska",
    "Southern Europe": "Southern Europe",
    "Northern Europe": "Northern Europe",
    "Transatlantic": None,          # no sell-side counterpart; left unmapped
}
# Sell-side regions this panel does not collect at all.
NOT_COLLECTED = ("Australia/New Zealand",)

BALCONY = "balcony"


def jefferies_region(region: str | None) -> str | None:
    return JEFFERIES_REGION.get(region or "")


def price_month(scrape_date: str) -> str:
    return scrape_date[:7]


def months_between(sail_month: str, price_month_: str) -> int:
    """Sail month minus price month, in whole months ("2027-01" - "2026-09" = 4)."""
    sy, sm = int(sail_month[:4]), int(sail_month[5:7])
    py, pm = int(price_month_[:4]), int(price_month_[5:7])
    return (sy - py) * 12 + (sm - pm)


# -- loading ----------------------------------------------------------------

def balcony_clause(tier: str, **filters: Any) -> tuple[str, list[Any]]:
    """The one filter every sheet here shares.

    Priced cells only -- an unpriced or sold-out cell has no price to take a
    minimum of -- cruise-only, balcony, and the peer-comparable exclusion. A
    `limited` cell would count as priced, but neither line publishes that
    state, so the inclusion is vacuous in practice.
    """
    clause, args = an._where(tier, priced_only=True, product="cruise_only",
                             **filters)
    clause += " AND cabin_category = ? AND availability_status != 'sold_out'"
    args = [*args, BALCONY]
    dropped = peer_excluded_subcategories()
    if dropped:
        clause += (" AND (cabin_subcategory IS NULL OR cabin_subcategory "
                   "NOT IN (%s))" % ",".join("?" * len(dropped)))
        args = [*args, *dropped]
    return clause, args


def load_balcony(conn: sqlite3.Connection, tier: str, **filters: Any
                 ) -> list[sqlite3.Row]:
    clause, args = balcony_clause(tier, **filters)
    return conn.execute(
        f"""SELECT line, region, sailing_id, cabin_subcategory, market,
                   vendor_category_code, ship, itinerary_code, nights,
                   sail_date, scrape_date, price_pppn, price_per_person,
                   substr(sail_date, 1, 7) AS sail_month
            FROM observations WHERE {clause}""", args).fetchall()


def _cell(r: sqlite3.Row) -> tuple:
    return (r["line"], r["sailing_id"], r["cabin_subcategory"], r["market"])


# -- sheet 1: cohort index on the Jefferies basis ---------------------------

def cohort_index_jefferies(conn: sqlite3.Connection, *, tier: str,
                           min_matched: int = 5) -> an.Result:
    """The Cohort index layout, rebuilt on MINIMUM balcony pppn.

    Two minima side by side, per (line, region, sail month) and scrape date:

    * index_naive_min -- min over ALL priced balcony cells that date. This is
      the sell-side-comparable number: whatever was listed.
    * index_matched_min -- min over cells priced on BOTH the base date and this
      date. A cheap cabin selling out cannot move it.

    mix_effect_min_pp = naive minus matched: the part of the minimum's move
    that is cabins leaving the book rather than prices changing.

    index_matched, index_naive and mix_effect_pp are the existing MEAN-based
    figures, computed exactly as on the Cohort index sheet but on this sheet's
    balcony-only universe, kept for reference; mean_pppn_* are the matched
    basket means behind them.
    """
    clause, args = balcony_clause(tier)
    basis = an._basis(conn, tier, clause, args)
    rows = load_balcony(conn, tier)
    if not rows:
        return an.Result("cohort_index_jefferies", basis, [],
                         notes=("no priced balcony rows in scope",))

    book: dict[tuple, dict[str, dict[tuple, float]]] = collections.defaultdict(
        lambda: collections.defaultdict(dict))
    for r in rows:
        book[(r["line"], r["region"], r["sail_month"])][r["scrape_date"]][
            _cell(r)] = r["price_pppn"]

    out: list[dict[str, Any]] = []
    for key in sorted(book, key=lambda k: tuple(map(str, k))):
        line, region, cohort = key
        dates = sorted(book[key])
        base = dates[0]
        base_cells = book[key][base]
        base_naive_min = min(base_cells.values())
        base_naive_mean = statistics.fmean(base_cells.values())
        base_d = date.fromisoformat(base)
        for d in dates:
            now = book[key][d]
            matched = set(base_cells) & set(now)
            m_base = [base_cells[k] for k in matched]
            m_now = [now[k] for k in matched]
            s_base, s_now = sum(m_base), sum(m_now)
            naive_min = min(now.values())
            naive_mean = statistics.fmean(now.values())
            rel = sorted((now[k] / base_cells[k] - 1) * 100
                         for k in matched if base_cells[k])
            idx_naive_min = round(100 * naive_min / base_naive_min, 2)
            idx_matched_min = (round(100 * min(m_now) / min(m_base), 2)
                               if matched else None)
            idx_matched = round(100 * s_now / s_base, 2) if s_base else None
            idx_naive = round(100 * naive_mean / base_naive_mean, 2)
            row = {
                "line": line, "region": region,
                "jefferies_region": jefferies_region(region),
                "cabin_category": BALCONY, "cohort": cohort,
                "base_date": base, "scrape_date": d,
                "price_month": price_month(d),
                "days_since_base": (date.fromisoformat(d) - base_d).days,
                "base_cells": len(base_cells), "cells_now": len(now),
                "matched_cells": len(matched),
                "index_naive_min": idx_naive_min,
                "index_matched_min": idx_matched_min,
                "mix_effect_min_pp": (round(idx_naive_min - idx_matched_min, 2)
                                      if idx_matched_min is not None else None),
                "naive_min_pppn_base": round(base_naive_min, 2),
                "naive_min_pppn_now": round(naive_min, 2),
                "basket_base_pppn": round(min(m_base), 2) if matched else None,
                "basket_now_pppn": round(min(m_now), 2) if matched else None,
                "mean_pppn_base": (round(statistics.fmean(m_base), 2)
                                   if matched else None),
                "mean_pppn_now": (round(statistics.fmean(m_now), 2)
                                  if matched else None),
                "index_matched": idx_matched,
                "median_cell_change_pct": (round(statistics.median(rel), 2)
                                           if rel else None),
                "index_naive": idx_naive,
                "mix_effect_pp": (round(idx_naive - idx_matched, 2)
                                  if idx_matched is not None else None),
                "attrition_pct": round(100 * (len(base_cells) - len(matched))
                                       / len(base_cells), 1),
                "entered_cells": len(set(now) - set(base_cells)),
            }
            n = len(matched)
            row["sample"] = ("BASE" if d == base
                             else "INSUFFICIENT" if n < min_matched
                             else "VERY THIN" if n < an.VERY_THIN_CELLS
                             else "THIN" if n < an.THIN_CELLS
                             else "ok") + f" (matched={n}, cells_now={len(now)})"
            out.append(row)

    basis = basis.with_caveat(an._history_note(list(basis.scrape_dates)))
    return an.Result("cohort_index_jefferies", basis, out, notes=(
        "Balcony only, cruise-only, NCL MINISUITE excluded as not "
        "peer-comparable, suites excluded. Priced cells only: a sold-out or "
        "unpriced cell has no price to take a minimum of.",
        "One priced balcony cell per sailing on both lines, so every minimum "
        "here is a minimum over SAILINGS in the sail month.",
        "index_naive_min is the sell-side-comparable number. index_matched_min "
        "holds the sailings fixed. mix_effect_min_pp = naive minus matched: the "
        "part of the minimum's move that is sailings leaving or entering.",
        "basket_*_pppn are MINIMA over the matched sailings (the existing sheet "
        "shows means there). naive_min_pppn_* are minima over all sailings. "
        "mean_pppn_* are the matched-basket MEANS, so mean/min can be formed.",
        "index_matched, index_naive and mix_effect_pp are the existing "
        "mean-based measures on this sheet's balcony universe, for reference.",
        "Base = the first scrape date on which that line x region x sail month "
        "was seen priced. Rows with different base_date are not comparable in "
        "levels of the index.",
    ))


# -- sheet 2: the grid ------------------------------------------------------

def min_provenance(conn: sqlite3.Connection, *, tier: str, line: str,
                   region: str, sail_month: str, price_month_: str
                   ) -> dict[str, Any] | None:
    """Which sailing and cabin produced the minimum in one grid cell."""
    clause, args = balcony_clause(tier, lines=[line], regions=[region])
    clause += " AND substr(sail_date, 1, 7) = ? AND substr(scrape_date, 1, 7) = ?"
    args = [*args, sail_month, price_month_]
    r = conn.execute(
        f"""SELECT line, region, ship, itinerary_code, sailing_id, sail_date,
                   nights, cabin_category, cabin_subcategory,
                   vendor_category_code, price_per_person, price_pppn,
                   scrape_date
            FROM observations WHERE {clause}
            ORDER BY price_pppn ASC, scrape_date ASC, sailing_id ASC
            LIMIT 1""", args).fetchone()
    return dict(r) if r else None


def jefferies_grid(conn: sqlite3.Connection, *, tier: str) -> an.Result:
    """Line x region x sail month down the side, price month across the top."""
    clause, args = balcony_clause(tier)
    basis = an._basis(conn, tier, clause, args)
    rows = load_balcony(conn, tier)
    if not rows:
        return an.Result("jefferies_grid", basis, [], notes=("no rows",))

    months = sorted({price_month(r["scrape_date"]) for r in rows})
    scrapes_in: dict[str, list[str]] = collections.defaultdict(list)
    for d in sorted({r["scrape_date"] for r in rows}):
        scrapes_in[price_month(d)].append(d)

    cell: dict[tuple, dict[str, dict[str, Any]]] = collections.defaultdict(dict)
    for r in rows:
        k = (r["line"], r["region"], r["sail_month"])
        pm = price_month(r["scrape_date"])
        cur = cell[k].get(pm)
        if cur is None:
            cur = cell[k][pm] = {"min": r["price_pppn"], "row": r,
                                 "sailings": set()}
        cur["sailings"].add(r["sailing_id"])
        # ties broken on the earliest scrape then sailing id, as in provenance
        if (r["price_pppn"], r["scrape_date"], r["sailing_id"]) < (
                cur["min"], cur["row"]["scrape_date"], cur["row"]["sailing_id"]):
            cur["min"], cur["row"] = r["price_pppn"], r

    out: list[dict[str, Any]] = []
    for k in sorted(cell, key=lambda x: tuple(map(str, x))):
        line, region, sail_month = k
        seen = [m for m in months if m in cell[k]]
        first = seen[0]
        first_min = cell[k][first]["min"]
        d = {"line": line, "jefferies_region": jefferies_region(region),
             "region": region, "sail_month": sail_month,
             "first_price_month": first}
        for m in months:
            c = cell[k].get(m)
            d[f"min_pppn {m}"] = round(c["min"], 2) if c else None
        for m in months:
            c = cell[k].get(m)
            d[f"chg_vs_first {m}"] = (round(100 * (c["min"] / first_min - 1), 2)
                                      if c else None)
        for m in months:
            d[f"months_to_sail {m}"] = months_between(sail_month, m)
        for m in months:
            c = cell[k].get(m)
            d[f"sailings {m}"] = len(c["sailings"]) if c else 0
        for m in months:
            c = cell[k].get(m)
            if c:
                w = c["row"]
                d[f"min_set_by {m}"] = (f"{w['ship']} {w['sail_date']} "
                                        f"{w['nights']}n obs {w['scrape_date']}")
            else:
                d[f"min_set_by {m}"] = None
        out.append(d)

    obs = "; ".join(f"{m}: {len(scrapes_in[m])} scrape(s) "
                    f"({', '.join(scrapes_in[m])})" for m in months)
    basis = basis.with_caveat(
        f"PRICE-MONTH COVERAGE: {obs}. A minimum over fewer observations can "
        "only be equal or higher than one over more, so a thinly observed "
        "price month is biased UP against a daily-observed series.")
    return an.Result("jefferies_grid", basis, out, notes=(
        "min_pppn = minimum balcony pppn over every sailing in the sail month "
        "and every scrape date in the price month. This is the sell-side "
        "construction.",
        "chg_vs_first = % change against the first price month observed for "
        "that row, not a fixed calendar month.",
        "months_to_sail = sail month minus price month.",
        "sailings = distinct sailings priced in that sail month during that "
        "price month -- the sample behind each minimum.",
        "min_set_by = the sailing that produced the minimum, and the scrape it "
        "was observed on. The minimum is a single cabin; this names it.",
    ))


# -- sheet 3: min vs mean by lead -------------------------------------------

def min_vs_mean_by_lead(conn: sqlite3.Connection, *, tier: str,
                        scrape_date: str | None = None) -> an.Result:
    """How far above the minimum the typical balcony fare sits, by lead time.

    The ratio is formed ACROSS SAILINGS within one line x region x sail month
    x days-to-departure band on one scrape date: mean over those sailings'
    balcony pppn divided by the minimum over them. That is the factor between
    a mean-based figure and the sell-side minimum for the same sail month.

    It is not formed per sailing, because each sailing carries one priced
    balcony cell, so a per-sailing mean/min is 1.0 by construction. Groups of
    a single sailing are likewise 1.0 by construction and are excluded and
    counted rather than allowed to pull the median to 1.
    """
    scrape_date = scrape_date or an._latest_scrape_date(conn, tier)
    clause, args = balcony_clause(tier, scrape_date=scrape_date)
    basis = an._basis(conn, tier, clause, args)
    rows = load_balcony(conn, tier, scrape_date=scrape_date)
    if not rows:
        return an.Result("min_vs_mean_by_lead", basis, [], notes=("no rows",))

    groups: dict[tuple, dict[str, float]] = collections.defaultdict(dict)
    for r in rows:
        dtd = (date.fromisoformat(r["sail_date"][:10])
               - date.fromisoformat(scrape_date)).days
        band = an.dtd_band(dtd)
        if band is None:
            continue
        groups[(r["line"], r["region"], band, r["sail_month"])][
            r["sailing_id"]] = r["price_pppn"]

    agg: dict[tuple, dict[str, Any]] = collections.defaultdict(
        lambda: {"ratios": [], "sailings": 0, "single": 0, "months": set()})
    for (line, region, band, month), prices in groups.items():
        a = agg[(line, region, band)]
        a["sailings"] += len(prices)
        a["months"].add(month)
        if len(prices) < 2:
            a["single"] += 1
            continue
        vals = list(prices.values())
        a["ratios"].append(statistics.fmean(vals) / min(vals))

    out: list[dict[str, Any]] = []
    for key in sorted(agg, key=lambda k: (k[0], str(k[1]), an._band_order(k[2]))):
        line, region, band = key
        a = agg[key]
        rs = sorted(a["ratios"])
        out.append({
            "line": line, "region": region,
            "jefferies_region": jefferies_region(region),
            "dtd_band": band,
            "median_mean_over_min": round(statistics.median(rs), 4) if rs else None,
            "p25_mean_over_min": round(rs[max(0, len(rs) // 4 - 1)], 4) if rs else None,
            "p75_mean_over_min": (round(rs[min(len(rs) - 1, 3 * len(rs) // 4)], 4)
                                  if rs else None),
            "sailings": a["sailings"],
            "sail_month_groups_used": len(rs),
            "single_sailing_groups_excluded": a["single"],
            "sail_months": ", ".join(sorted(a["months"])),
        })
    return an.Result("min_vs_mean_by_lead", basis, out, notes=(
        f"Scrape date {scrape_date}. Band edges are the Booking curve's: "
        f"{list(an.DTD_EDGES)}.",
        "Ratio = mean over sailings / minimum over sailings, within one sail "
        "month and band. 1.30 means the average balcony fare sits 30% above "
        "the cheapest one in that sail month.",
        "NOT a per-sailing ratio: each sailing has one priced balcony cell, so "
        "that ratio is 1.0 by construction. Single-sailing groups are 1.0 for "
        "the same reason and are excluded and counted.",
        "sailings counts every sailing in the band, including those in "
        "excluded single-sailing groups.",
    ))


# -- index-sheet facts ------------------------------------------------------

def basis_facts(conn: sqlite3.Connection, tier: str, scrape_date: str
                ) -> list[tuple[str, str]]:
    """Plain-words documentation of the basis, measured from the data."""
    def one(sql: str, args: Sequence[Any] = ()) -> Any:
        return conn.execute(sql, list(args)).fetchone()[0]

    W = "tier = ? AND scrape_date = ?"
    wa = [tier, scrape_date]
    priced = one(f"SELECT COUNT(*) FROM observations WHERE {W} "
                 "AND price_pppn IS NOT NULL", wa)
    taxed = one(f"SELECT COUNT(*) FROM observations WHERE {W} "
                "AND price_pppn IS NOT NULL AND taxes_fees IS NOT NULL", wa)
    dbl = one(f"SELECT COUNT(*) FROM observations WHERE {W} "
              "AND price_pppn IS NOT NULL "
              "AND ABS(price_total - 2 * price_per_person) < 0.01", wa)
    pppn_ok = one(f"SELECT COUNT(*) FROM observations WHERE {W} "
                  "AND price_pppn IS NOT NULL AND nights > 0 "
                  "AND ABS(price_pppn - price_per_person * 1.0 / nights) < 0.01", wa)
    bases = ", ".join(f"{r[0]}: {r[1]}" for r in conn.execute(
        f"SELECT line, GROUP_CONCAT(DISTINCT price_basis) FROM observations "
        f"WHERE {W} AND price_pppn IS NOT NULL GROUP BY line", wa))
    pkg = conn.execute(
        f"""SELECT line, COUNT(*) FROM observations WHERE {W}
            AND cabin_category = 'balcony' AND is_package = 1
            AND price_pppn IS NOT NULL GROUP BY line""", wa).fetchall()
    pkg_txt = (", ".join(f"{r[0]}: {r[1]}" for r in pkg)
               or "none") + f" (priced balcony package rows on {scrape_date})"
    states = ", ".join(f"{r[0]}: {r[1]}" for r in conn.execute(
        f"""SELECT line, GROUP_CONCAT(DISTINCT availability_status)
            FROM observations WHERE {W} GROUP BY line""", wa))
    dropped = peer_excluded_subcategories()
    mini_min = one(
        f"""WITH g AS (
              SELECT line, region, substr(sail_date,1,7) sm, scrape_date,
                     MIN(CASE WHEN cabin_subcategory NOT IN ({','.join('?'*len(dropped))})
                              THEN price_pppn END) AS kept,
                     MIN(price_pppn) AS withall
              FROM observations
              WHERE tier = ? AND cabin_category = 'balcony' AND is_package = 0
                AND price_pppn IS NOT NULL
              GROUP BY 1,2,3,4)
            SELECT COUNT(*) FROM g WHERE withall < kept""",
        [*dropped, tier]) if dropped else 0
    scrapes = [r[0] for r in conn.execute(
        "SELECT DISTINCT scrape_date FROM observations WHERE tier = ? "
        "ORDER BY 1", [tier])]
    window = conn.execute(
        "SELECT MIN(sail_date), MAX(sail_date) FROM observations WHERE tier = ?",
        [tier]).fetchone()

    return [
        ("Price field",
         f"The booking API's published fare per cabin category per sailing "
         f"({bases}). NCL's combinedPrice includes its More at Sea bundle; "
         f"Carnival's rooms.price does not bundle equivalents."),
        ("Taxes, fees, port charges",
         f"EXCLUDED. taxes_fees is populated on {taxed} of {priced} priced rows "
         f"on {scrape_date}. NCL's API carries no tax field and its disclaimer "
         f"states taxes, fees and port expenses are additional. Carnival's tax "
         f"fields exist but read 0 on every room and every itinerary "
         f"(minTaxesAndFees: 0), so the amount cannot be recovered. The "
         f"sell-side series includes taxes; ours sit below it by the per-night "
         f"tax amount, which this panel cannot measure."),
        ("Per person / per night",
         f"price_pppn = price_per_person / nights, exactly, on {pppn_ok} of "
         f"{priced} priced rows. Taxes excluded, as above."),
        ("Occupancy",
         f"Double occupancy: price_total = 2 x price_per_person on {dbl} of "
         f"{priced} priced rows. pppn is one guest's share."),
        ("Cruise-only vs packages",
         f"Cruise-only kept, as in the panel. Excluded: {pkg_txt}. Land+cruise "
         f"packages carry a package fare over cruise nights, which would "
         f"overstate pppn."),
        ("Cabin",
         f"Balcony only. Suites excluded. {', '.join(dropped) or 'nothing'} "
         f"dropped from NCL's balcony bucket as not peer-comparable. Dropping "
         f"it changed the minimum in {mini_min} line x region x sail month x "
         f"scrape groups -- it is the dearer grade, so it almost never sets "
         f"the minimum."),
        ("Availability", f"Priced cells only. States present: {states}. No line "
         "publishes a 'limited' state, so 'include limited' is vacuous; "
         "sold-out and solo-only cells carry no price and drop out."),
        ("Observation frequency",
         f"This tier ({tier}) is observed on {len(scrapes)} dates: "
         f"{', '.join(scrapes)}. A minimum over fewer observations in a price "
         f"month is equal to or higher than one over more; daily-marker "
         f"observations are not pooled in, to keep tiers separate."),
        ("Sail window", f"{window[0]} to {window[1]}. Sail months before the "
         "window start are absent; the sell-side covers whatever was listed."),
        ("Itinerary length", "No filter, matching the sell-side method."),
        ("Brands", "Norwegian Cruise Line and Carnival Cruise Line brands only; "
         "US market, USD, verified on every run."),
    ]


def region_table(conn: sqlite3.Connection, tier: str) -> list[dict[str, Any]]:
    counts = dict(conn.execute(
        "SELECT region, COUNT(DISTINCT sailing_id) FROM observations "
        "WHERE tier = ? GROUP BY region", [tier]).fetchall())
    out = []
    for ours, theirs in JEFFERIES_REGION.items():
        note = ""
        if ours == "Caribbean":
            note = ("Includes Bahamas on both lines: NCL maps BAHAMAS and "
                    "Carnival maps region code BH to Caribbean. Matches the "
                    "sell-side's combined Caribbean/Bahamas.")
        if theirs is None:
            note = "No sell-side counterpart; left unmapped."
        out.append({"our_region": ours, "jefferies_region": theirs or "(unmapped)",
                    "sailings_collected": counts.get(ours, 0), "note": note})
    for theirs in NOT_COLLECTED:
        out.append({"our_region": "(none)", "jefferies_region": theirs,
                    "sailings_collected": 0, "note": "not collected"})
    return out


# -- workbook ---------------------------------------------------------------

SHEETS = (
    ("Cohort index (Jefferies basis)", "time series (all scrape dates)",
     "The Cohort index layout on MINIMUM balcony pppn: naive minimum "
     "(sell-side comparable) beside the matched-sailing minimum."),
    ("Jefferies grid", "time series (all scrape dates, by price month)",
     "Minimum balcony pppn by line x region x sail month (rows) and price "
     "month (columns), with % change and months to sail."),
    ("Min vs mean by lead", "latest scrape date",
     "Mean over sailings / minimum over sailings within a sail month, by "
     "days-to-departure band."),
)


def build(conn: sqlite3.Connection, *, tier: str, scrape_date: str):
    return [cohort_index_jefferies(conn, tier=tier),
            jefferies_grid(conn, tier=tier),
            min_vs_mean_by_lead(conn, tier=tier, scrape_date=scrape_date)]


def _digest(results: Sequence[an.Result]) -> tuple[str, dict[str, Any]]:
    body = json.dumps([r.to_dict() for r in results], sort_keys=True,
                      default=str).encode("utf-8")
    return hashlib.sha256(body).hexdigest(), {
        "digest": hashlib.sha256(body).hexdigest(),
        "sheets": {name: {"rows": len(r.rows), "status": "ok"}
                   for (name, _, _), r in zip(SHEETS, results)}}


def write_workbook(results: Sequence[an.Result], path: str, *, tier: str,
                   scrape_date: str, facts: Sequence[tuple[str, str]],
                   regions: Sequence[dict[str, Any]]) -> None:
    from openpyxl import Workbook

    st = rp._sheet_styles()
    wb = Workbook()
    idx = wb.active
    idx.title = "Index"
    r = 1
    idx.cell(row=r, column=1, value=f"Jefferies-basis workbook -- {tier}, "
             f"latest scrape {scrape_date}").font = st["title"]
    r += 1
    idx.cell(row=r, column=1, value=(
        "Rebuilt on the sell-side's MINIMUM balcony price so the two series "
        "can be compared. The panel's own workbook is unchanged. Regenerate "
        "with `python -m panel.report_jefferies`.")).font = st["small"]
    r += 2
    idx.cell(row=r, column=1, value="Basis, in plain words").font = st["head"]
    r += 1
    for key, value in facts:
        idx.cell(row=r, column=1, value=key).font = st["head"]
        c = idx.cell(row=r, column=2, value=value)
        c.alignment = st["wrap"]
        r += 1
    idx.cell(row=r, column=1, value="Generated at (UTC)").font = st["head"]
    idx.cell(row=r, column=2, value=datetime.now(timezone.utc)
             .replace(microsecond=0).isoformat())
    r += 2
    idx.cell(row=r, column=1, value="Region mapping").font = st["head"]
    r += 2
    r = rp._write_table(idx, list(regions), r, st) + 1
    idx.cell(row=r, column=1, value="Sheets").font = st["head"]
    r += 2
    rp._write_table(idx, [{"sheet": n, "time scope": t, "rows": len(res.rows),
                           "what it shows": p}
                          for (n, t, p), res in zip(SHEETS, results)], r, st)
    idx.freeze_panes = None
    idx.column_dimensions["A"].width = 30
    idx.column_dimensions["B"].width = 110

    for (name, scope, purpose), res in zip(SHEETS, results):
        ws = wb.create_sheet(name[:31])
        row = 1
        ws.cell(row=row, column=1, value=name).font = st["title"]
        row += 1
        ws.cell(row=row, column=1, value=purpose).font = st["small"]
        row += 1
        ws.cell(row=row, column=1, value=f"Time scope: {scope}.").font = st["small"]
        row += 1
        ws.cell(row=row, column=1, value=res.basis.label()).font = st["small"]
        row += 2
        for cav in res.basis.caveats:
            c = ws.cell(row=row, column=1, value="! " + cav)
            c.font, c.fill = st["warn"], st["warnfill"]
            row += 1
        for note in res.notes:
            ws.cell(row=row, column=1, value="- " + note).font = st["small"]
            row += 1
        row += 1
        rp._write_table(ws, res.rows, row, st)
    wb.save(path)


def report_name(tier: str, scrape_date: str) -> str:
    return f"{scrape_date}__{tier}__jefferies-basis.xlsx"


def write_report(db_path: str, *, tier: str = "weekly-full",
                 scrape_date: str | None = None,
                 out_dir: str = rp.REPORT_DIR,
                 reason: str = "manual run", force: bool = False
                 ) -> dict[str, Any]:
    """Build the workbook. Unchanged content is not rewritten; changed content
    is logged to REGENERATIONS.jsonl, the same rule the panel's workbook obeys."""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        scrape_date = scrape_date or rp.latest_scrape_date(conn, tier)
        if scrape_date is None:
            return {"path": None, "written": False,
                    "note": f"no observations in tier {tier!r}"}
        results = build(conn, tier=tier, scrape_date=scrape_date)
        facts = basis_facts(conn, tier, scrape_date)
        regions = region_table(conn, tier)
    finally:
        conn.close()

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, report_name(tier, scrape_date))
    manifest_path = path.replace(".xlsx", ".manifest.json")
    digest, manifest = _digest(results)
    before = None
    if os.path.exists(manifest_path):
        try:
            with open(manifest_path, encoding="utf-8") as fh:
                before = json.load(fh)
        except (OSError, ValueError):
            before = None
    existed = os.path.exists(path)
    summary = {"path": path, "tier": tier, "scrape_date": scrape_date,
               "sheets": manifest["sheets"], "written": False,
               "regenerated": False}
    if existed and (before or {}).get("digest") == digest and not force:
        summary["note"] = "content identical to the existing file; not rewritten"
        return summary

    write_workbook(results, path, tier=tier, scrape_date=scrape_date,
                   facts=facts, regions=regions)
    manifest.update({"tier": tier, "scrape_date": scrape_date,
                     "basis": "jefferies",
                     "generated_at_utc": datetime.now(timezone.utc)
                     .replace(microsecond=0).isoformat()})
    with open(manifest_path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(manifest, fh, indent=1, sort_keys=True)
        fh.write("\n")
    summary["written"] = True
    if existed:
        summary["regenerated"] = True
        rp.log_regeneration(out_dir, report=os.path.basename(path),
                            reason=reason, tool="panel.report_jefferies",
                            before=before, after=manifest)
    return summary


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--tier", default="weekly-full", choices=list(an.TIERS))
    ap.add_argument("--db", default=os.path.join("data", "panel.sqlite"))
    ap.add_argument("--scrape-date", default=None,
                    help="default: the latest scrape date in the tier")
    ap.add_argument("--out-dir", default=rp.REPORT_DIR)
    ap.add_argument("--reason", default="manual run")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args(argv)
    if not os.path.exists(args.db):
        print(f"ERROR: no database at {args.db}", file=sys.stderr)
        return 2
    s = write_report(args.db, tier=args.tier, scrape_date=args.scrape_date,
                     out_dir=args.out_dir, reason=args.reason, force=args.force)
    if not s.get("path"):
        print(s.get("note"))
        return 0
    for name, meta in s["sheets"].items():
        print(f"  {name:<34} {meta['rows']:>6} rows")
    print(("wrote " if s["written"] else "unchanged: ") + s["path"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
