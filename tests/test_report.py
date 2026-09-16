"""The post-run report pipeline.

Three properties carry the weight here.

1. A scheduled run that nobody watches must produce a workbook whatever one
   analysis does, and must say loudly when one of them failed. Silence is the
   failure mode that matters: this thing runs at 06:10 UTC.
2. A time series across a change in the COLLECTOR's own scope is not a time
   series. The suite must refuse to present one as if it were.
3. Regenerating a committed workbook is allowed -- it is derived -- but never
   silent. Unchanged content must not rewrite the file at all, and changed
   content must leave a plain-text explanation next to the binary diff.
"""
import itertools
import json
import os
import sqlite3

import openpyxl
import pytest

from panel import report as rp
from panel.schema import DDL

NCL = "Norwegian Cruise Line"
CCL = "Carnival Cruise Line"

_ids = itertools.count(1)


def make_row(**kw):
    row = {
        "scrape_ts_utc": "2026-09-15T18:00:00+00:00",
        "scrape_date": "2026-09-15",
        "line": NCL, "brand": "NCL", "ship": "Norwegian Joy",
        "sailing_id": str(next(_ids)), "itinerary_code": "X1",
        "sail_date": "2027-02-01", "nights": 7,
        "is_package": 0, "itinerary_nights": 7,
        "region": "Caribbean", "market": "US", "currency": "USD",
        "cabin_category": "balcony", "cabin_subcategory": "BALCONY",
        "price_total": 2800.0, "price_pppn": 200.0,
        "availability_status": "available", "tier": "weekly-full",
    }
    row.update(kw)
    return row


def write(path, rows):
    conn = sqlite3.connect(path)
    conn.executescript(DDL)
    for r in rows:
        conn.execute(f"INSERT INTO observations ({', '.join(r)}) "
                     f"VALUES ({', '.join('?' * len(r))})", list(r.values()))
    conn.commit()
    conn.close()
    return path


@pytest.fixture
def db(tmp_path):
    """Two lines, two regions, one collection date."""
    rows = []
    for line in (NCL, CCL):
        for region in ("Caribbean", "Alaska"):
            for cat in ("inside", "balcony", "suite"):
                for i in range(6):
                    rows.append(make_row(
                        line=line, region=region, cabin_category=cat,
                        cabin_subcategory=cat.upper(),
                        sail_date=f"2027-02-{i + 1:02d}",
                        price_pppn=100.0 + i * 10))
    return write(str(tmp_path / "panel.sqlite"), rows)


@pytest.fixture
def config(tmp_path):
    path = tmp_path / "panel.yaml"
    path.write_text("""
storage: {db_path: data/panel.sqlite, raw_archive_path: data/raw}
http:
  user_agent: t
  min_interval_s: 2.0
  max_retries: 4
  backoff_base_s: 2.0
  backoff_cap_s: 60.0
  timeout_s: 60.0
  obey_robots: true
lines:
  ncl:
    line: "Norwegian Cruise Line"
    base_url: "https://example.invalid"
    regions: [Caribbean, Alaska]
  carnival:
    line: "Carnival Cruise Line"
    base_url: "https://example.invalid"
    regions: [Caribbean, Transatlantic]
  royal:
    line: "Royal Caribbean"
    enabled: false
    base_url: "https://example.invalid"
    regions: [Antarctica]
tiers:
  weekly-full:
    sail_window: {start: "2026-10-01", end: "2027-08-31"}
  daily-marker:
    sail_window: {start: "2026-10-01", end: "2026-12-31"}
    marker_only: true
""", encoding="utf-8")
    return str(path)


def conn_for(db):
    c = sqlite3.connect(db)
    c.row_factory = sqlite3.Row
    return c


# -- region scope -----------------------------------------------------------

class TestSuiteRegions:
    def test_union_of_enabled_lines_in_config_order(self, config):
        assert rp.suite_regions(config) == ["Caribbean", "Alaska", "Transatlantic"]

    def test_disabled_line_contributes_nothing(self, config):
        assert "Antarctica" not in rp.suite_regions(config)

    def test_reads_the_real_config_without_raising(self):
        # The list is derived, never hardcoded -- that is the whole point.
        assert "Caribbean" in rp.suite_regions("config/panel.yaml")


# -- running the suite ------------------------------------------------------

class TestRunSuite:
    def test_every_analysis_in_the_suite_runs(self, db):
        sheets = rp.run_suite(conn_for(db), tier="weekly-full",
                              scrape_date="2026-09-15",
                              regions=["Caribbean", "Alaska"])
        assert [s.spec.key for s in sheets] == [s.key for s in rp.SUITE]
        assert all(s.error is None for s in sheets)

    def test_one_failing_analysis_does_not_abort_the_rest(self, db, monkeypatch):
        def boom(*a, **k):
            raise RuntimeError("synthetic")

        monkeypatch.setattr(
            rp, "SUITE",
            tuple(rp.SheetSpec(s.key, s.sheet, s.scope,
                               boom if s.key == "depletion" else s.fn, s.purpose)
                  for s in rp.SUITE))
        sheets = rp.run_suite(conn_for(db), tier="weekly-full",
                              scrape_date="2026-09-15", regions=["Caribbean"])
        failed = [s for s in sheets if s.error]
        assert len(failed) == 1
        assert "synthetic" in failed[0].error
        assert failed[0].status.startswith("FAILED")
        assert sum(1 for s in sheets if s.result) == len(rp.SUITE) - 1

    def test_cross_section_is_pinned_to_the_date_time_series_is_not(self, db):
        sheets = {s.spec.key: s for s in rp.run_suite(
            conn_for(db), tier="weekly-full", scrape_date="2026-09-15",
            regions=["Caribbean", "Alaska"])}
        assert sheets["availability"].spec.scope is rp.CROSS_SECTION
        assert sheets["cohort-index"].spec.scope is rp.TIME_SERIES


# -- scope drift ------------------------------------------------------------

class TestScopeDrift:
    def test_no_drift_with_a_single_date(self, db):
        assert rp.scope_drift(conn_for(db), "weekly-full") == []

    def test_stable_scope_is_comparable(self, tmp_path):
        rows = [make_row(scrape_date=d, sailing_id=f"{d}-{i}",
                         sail_date="2027-02-01")
                for d in ("2026-09-15", "2026-09-22") for i in range(20)]
        drift = rp.scope_drift(conn_for(write(str(tmp_path / "p.sqlite"), rows)),
                               "weekly-full")
        assert [d["comparable"] for d in drift] == ["yes"]

    def test_widened_sail_window_is_flagged(self, tmp_path):
        rows = [make_row(scrape_date="2026-09-15", sail_date="2027-02-01")
                for _ in range(20)]
        rows += [make_row(scrape_date="2026-09-22",
                          sail_date="2026-10-05" if i else "2027-02-01")
                 for i in range(20)]
        drift = rp.scope_drift(conn_for(write(str(tmp_path / "p.sqlite"), rows)),
                               "weekly-full")
        assert drift[0]["comparable"] == "NO"
        assert "sail window moved" in drift[0]["why_not"]

    def test_new_region_is_flagged(self, tmp_path):
        rows = [make_row(scrape_date="2026-09-15") for _ in range(20)]
        rows += [make_row(scrape_date="2026-09-22") for _ in range(19)]
        rows += [make_row(scrape_date="2026-09-22", region="Transatlantic")]
        drift = rp.scope_drift(conn_for(write(str(tmp_path / "p.sqlite"), rows)),
                               "weekly-full")
        assert "regions added: Transatlantic" in drift[0]["why_not"]

    def test_a_line_appearing_is_flagged(self, tmp_path):
        rows = [make_row(scrape_date="2026-09-15") for _ in range(20)]
        rows += [make_row(scrape_date="2026-09-22") for _ in range(19)]
        rows += [make_row(scrape_date="2026-09-22", line=CCL)]
        drift = rp.scope_drift(conn_for(write(str(tmp_path / "p.sqlite"), rows)),
                               "weekly-full")
        assert "lines changed" in drift[0]["why_not"]

    def test_large_sailing_count_jump_is_flagged(self, tmp_path):
        rows = [make_row(scrape_date="2026-09-15") for _ in range(10)]
        rows += [make_row(scrape_date="2026-09-22") for _ in range(40)]
        drift = rp.scope_drift(conn_for(write(str(tmp_path / "p.sqlite"), rows)),
                               "weekly-full")
        assert drift[0]["comparable"] == "NO"
        assert "sailing count moved" in drift[0]["why_not"]
        assert drift[0]["sailings_delta_pct"] == 300.0

    def test_small_movement_is_not_flagged(self, tmp_path):
        rows = [make_row(scrape_date="2026-09-15") for _ in range(100)]
        rows += [make_row(scrape_date="2026-09-22") for _ in range(105)]
        drift = rp.scope_drift(conn_for(write(str(tmp_path / "p.sqlite"), rows)),
                               "weekly-full")
        assert drift[0]["comparable"] == "yes"

    def test_caveat_is_none_when_everything_is_comparable(self):
        assert rp.drift_caveat([{"comparable": "yes", "from_date": "a",
                                 "to_date": "b", "why_not": ""}]) is None

    def test_caveat_names_the_dates_and_the_reason(self):
        note = rp.drift_caveat([{"comparable": "NO", "from_date": "2026-09-15",
                                 "to_date": "2026-09-22",
                                 "why_not": "sail window moved"}])
        assert note.startswith("COLLECTION SCOPE CHANGED")
        assert "2026-09-15->2026-09-22" in note
        assert "sail window moved" in note


class TestDriftReachesTheSheets:
    @pytest.fixture
    def drifted(self, tmp_path):
        rows = [make_row(scrape_date="2026-09-15") for _ in range(10)]
        rows += [make_row(scrape_date="2026-09-22") for _ in range(40)]
        return write(str(tmp_path / "p.sqlite"), rows)

    def test_time_series_sheets_carry_the_warning(self, drifted):
        note = rp.drift_caveat(rp.scope_drift(conn_for(drifted), "weekly-full"))
        sheets = rp.run_suite(conn_for(drifted), tier="weekly-full",
                              scrape_date="2026-09-22", regions=["Caribbean"],
                              drift_note=note)
        for s in sheets:
            if s.spec.scope is rp.TIME_SERIES:
                assert note in s.result.basis.caveats, s.spec.key

    def test_cross_sectional_sheets_do_not(self, drifted):
        """One date is one date. The warning is about spanning dates, and a
        caveat that fires everywhere is a caveat nobody reads."""
        note = rp.drift_caveat(rp.scope_drift(conn_for(drifted), "weekly-full"))
        sheets = rp.run_suite(conn_for(drifted), tier="weekly-full",
                              scrape_date="2026-09-22", regions=["Caribbean"],
                              drift_note=note)
        for s in sheets:
            if s.spec.scope is rp.CROSS_SECTION:
                assert note not in s.result.basis.caveats, s.spec.key

    def test_status_column_shows_it(self, drifted):
        note = rp.drift_caveat(rp.scope_drift(conn_for(drifted), "weekly-full"))
        sheets = rp.run_suite(conn_for(drifted), tier="weekly-full",
                              scrape_date="2026-09-22", regions=["Caribbean"],
                              drift_note=note)
        ts = [s for s in sheets if s.spec.scope is rp.TIME_SERIES and s.rows]
        assert ts and all(s.status == "COLLECTION SCOPE CHANGED BETWEEN DATES"
                          for s in ts)


# -- region coverage --------------------------------------------------------

class TestRegionCoverage:
    def test_a_configured_region_with_no_rows_is_still_listed(self, db):
        cov = rp.region_coverage(conn_for(db), "weekly-full", "2026-09-15",
                                 ["Caribbean", "Alaska", "Bermuda"])
        bermuda = [r for r in cov if r["region"] == "Bermuda"][0]
        assert bermuda["observations"] == 0
        assert "absent from every sheet" in bermuda["note"]

    def test_a_region_in_the_data_but_not_the_config_is_flagged(self, db):
        cov = rp.region_coverage(conn_for(db), "weekly-full", "2026-09-15",
                                 ["Caribbean"])
        alaska = [r for r in cov if r["region"] == "Alaska"][0]
        assert "NOT IN CONFIG" in alaska["note"]


# -- digest and regeneration ledger -----------------------------------------

class TestDigest:
    def test_same_results_same_digest(self, db):
        kw = dict(tier="weekly-full", scrape_date="2026-09-15",
                  regions=["Caribbean"])
        a, _ = rp.content_digest(rp.run_suite(conn_for(db), **kw))
        b, _ = rp.content_digest(rp.run_suite(conn_for(db), **kw))
        assert a == b

    def test_different_scope_different_digest(self, db):
        a, _ = rp.content_digest(rp.run_suite(
            conn_for(db), tier="weekly-full", scrape_date="2026-09-15",
            regions=["Caribbean"]))
        b, _ = rp.content_digest(rp.run_suite(
            conn_for(db), tier="weekly-full", scrape_date="2026-09-15",
            regions=["Caribbean", "Alaska"]))
        assert a != b


class TestWriteReport:
    def test_writes_workbook_and_manifest(self, db, config, tmp_path):
        out = str(tmp_path / "reports")
        s = rp.write_report(db, tier="weekly-full", out_dir=out,
                            config_path=config)
        assert s["written"] and not s["regenerated"]
        assert os.path.exists(s["path"])
        assert os.path.exists(s["path"].replace(".xlsx", ".manifest.json"))

    def test_unchanged_rerun_does_not_touch_the_file(self, db, config, tmp_path):
        out = str(tmp_path / "reports")
        first = rp.write_report(db, tier="weekly-full", out_dir=out,
                                config_path=config)
        mtime = os.path.getmtime(first["path"])
        again = rp.write_report(db, tier="weekly-full", out_dir=out,
                                config_path=config)
        assert not again["written"]
        assert os.path.getmtime(first["path"]) == mtime
        assert not os.path.exists(os.path.join(out, rp.REGEN_LEDGER))

    def test_changed_content_is_written_and_logged(self, db, config, tmp_path):
        out = str(tmp_path / "reports")
        rp.write_report(db, tier="weekly-full", out_dir=out, config_path=config)

        conn = sqlite3.connect(db)
        conn.execute("UPDATE observations SET availability_status = 'sold_out', "
                     "price_total = NULL, price_pppn = NULL "
                     "WHERE cabin_category = 'suite'")
        conn.commit()
        conn.close()

        again = rp.write_report(db, tier="weekly-full", out_dir=out,
                                config_path=config,
                                reason="suites went sold out")
        assert again["written"] and again["regenerated"]
        entries = [json.loads(l) for l in
                   open(again["ledger"], encoding="utf-8") if l.strip()]
        assert len(entries) == 1
        assert entries[0]["reason"] == "suites went sold out"
        assert entries[0]["digest_before"] != entries[0]["digest_after"]
        assert entries[0]["sheets_changed"], "a changed report must say what moved"

    def test_force_rewrites_identical_content(self, db, config, tmp_path):
        out = str(tmp_path / "reports")
        rp.write_report(db, tier="weekly-full", out_dir=out, config_path=config)
        again = rp.write_report(db, tier="weekly-full", out_dir=out,
                                config_path=config, force=True)
        assert again["written"] and again["regenerated"]

    def test_empty_tier_reports_no_scrape_date(self, tmp_path, config):
        empty = write(str(tmp_path / "empty.sqlite"), [])
        s = rp.write_report(empty, tier="daily-marker",
                            out_dir=str(tmp_path / "r"), config_path=config)
        assert s["scrape_date"] is None


# -- the workbook itself ----------------------------------------------------

class TestWorkbook:
    @pytest.fixture
    def wb(self, db, config, tmp_path):
        s = rp.write_report(db, tier="weekly-full",
                            out_dir=str(tmp_path / "reports"), config_path=config)
        return openpyxl.load_workbook(s["path"]), s

    def test_one_sheet_per_analysis_plus_an_index(self, wb):
        book, _ = wb
        assert book.sheetnames[0] == "Index"
        assert set(book.sheetnames[1:]) == {s.sheet for s in rp.SUITE}

    def test_index_states_the_currency_and_market(self, wb):
        book, _ = wb
        text = "\n".join(str(c.value) for row in book["Index"].iter_rows()
                         for c in row if c.value is not None)
        assert "USD" in text and "Currencies present" in text

    def test_index_declares_each_sheets_time_scope(self, wb):
        book, _ = wb
        text = "\n".join(str(c.value) for row in book["Index"].iter_rows()
                         for c in row if c.value is not None)
        assert rp.CROSS_SECTION in text and rp.TIME_SERIES in text

    def test_each_sheet_carries_its_basis_line(self, wb):
        book, _ = wb
        for spec in rp.SUITE:
            text = "\n".join(str(c.value) for row in book[spec.sheet].iter_rows()
                             for c in row if c.value is not None)
            assert "[weekly-full]" in text, spec.sheet

    def test_data_sheets_freeze_their_header(self, wb):
        book, _ = wb
        assert book["Availability"].freeze_panes is not None

    def test_prices_are_numbers_not_strings(self, wb):
        """A price written as text cannot be charted or averaged in Excel."""
        book, _ = wb
        ws = book["Peer gap"]
        header = None
        for row in ws.iter_rows():
            if any(c.value == "treatment_median_pppn" for c in row):
                header = {c.value: c.column for c in row}
                break
        assert header, "peer gap should have produced a table"
        col = header["treatment_median_pppn"]
        vals = [ws.cell(row=r, column=col).value
                for r in range(header and 1 or 1, ws.max_row + 1)]
        assert any(isinstance(v, (int, float)) for v in vals)

    def test_a_failed_analysis_is_visible_in_the_workbook(
            self, db, config, tmp_path, monkeypatch):
        def boom(*a, **k):
            raise RuntimeError("synthetic failure")

        monkeypatch.setattr(
            rp, "SUITE",
            tuple(rp.SheetSpec(s.key, s.sheet, s.scope,
                               boom if s.key == "promo-diff" else s.fn, s.purpose)
                  for s in rp.SUITE))
        s = rp.write_report(db, tier="weekly-full",
                            out_dir=str(tmp_path / "r"), config_path=config)
        assert s["failed"] == ["promo-diff"]
        book = openpyxl.load_workbook(s["path"])
        text = "\n".join(str(c.value) for row in book["Promo diff"].iter_rows()
                         for c in row if c.value is not None)
        assert "THIS ANALYSIS FAILED" in text and "synthetic failure" in text


# -- CLI --------------------------------------------------------------------

class TestCli:
    def test_runs_and_writes(self, db, config, tmp_path, capsys):
        rc = rp.main(["--tier", "weekly-full", "--db", db,
                      "--out-dir", str(tmp_path / "r"), "--config", config])
        assert rc == 0
        assert "availability" in capsys.readouterr().out

    def test_missing_database_exits_two(self, tmp_path):
        assert rp.main(["--tier", "weekly-full",
                        "--db", str(tmp_path / "nope.sqlite")]) == 2

    def test_analysis_failure_can_fail_the_run(self, db, config, tmp_path,
                                               monkeypatch):
        def boom(*a, **k):
            raise RuntimeError("synthetic")

        monkeypatch.setattr(
            rp, "SUITE",
            tuple(rp.SheetSpec(s.key, s.sheet, s.scope,
                               boom if s.key == "depletion" else s.fn, s.purpose)
                  for s in rp.SUITE))
        args = ["--tier", "weekly-full", "--db", db, "--config", config,
                "--out-dir", str(tmp_path / "r")]
        assert rp.main(args) == 0                        # still writes a report
        assert rp.main(args + ["--fail-on-analysis-error", "--force"]) == 1

    def test_github_summary_is_appended_when_asked(self, db, config, tmp_path,
                                                   monkeypatch):
        dest = tmp_path / "summary.md"
        monkeypatch.setenv("GITHUB_STEP_SUMMARY", str(dest))
        rp.main(["--tier", "weekly-full", "--db", db, "--config", config,
                 "--out-dir", str(tmp_path / "r"), "--github-summary"])
        assert "| sheet | rows | status |" in dest.read_text(encoding="utf-8")

    def test_drift_is_shouted_in_the_summary(self, tmp_path, config, capsys):
        rows = [make_row(scrape_date="2026-09-15") for _ in range(10)]
        rows += [make_row(scrape_date="2026-09-22") for _ in range(40)]
        path = write(str(tmp_path / "p.sqlite"), rows)
        rp.main(["--tier", "weekly-full", "--db", path, "--config", config,
                 "--out-dir", str(tmp_path / "r")])
        assert "SCOPE DRIFT" in capsys.readouterr().out
