"""The Jefferies-basis workbook: a minimum over sail month x price month.

The sell-side construction is a MINIMUM over every sailing in a sail month and
every observation in a price month. These tests pin that aggregation, the two
things it is confused by (sailings entering and sailings selling out), which
the matched-basket minimum is there to separate, and the filters that decide
what counts as a comparable balcony cabin.

Dates are synthetic and deliberately nowhere near the real collection
calendar: nothing here may depend on when the panel happened to run.
"""
import itertools
import os
import sqlite3

import openpyxl
import pytest

from panel import report_jefferies as rj
from panel.schema import DDL

NCL = "Norwegian Cruise Line"
CCL = "Carnival Cruise Line"
_ids = itertools.count(1)


def cell(*, sailing, scrape, price, line=NCL, region="Caribbean",
         sail_date="2031-03-10", nights=7, category="balcony", sub=None,
         status="available", package=0, ship="Ship A"):
    return {
        "scrape_ts_utc": scrape + "T06:00:00+00:00", "scrape_date": scrape,
        "line": line, "brand": "b", "ship": ship, "sailing_id": sailing,
        "itinerary_code": "IT" + sailing, "sail_date": sail_date,
        "nights": nights, "is_package": package, "region": region,
        "market": "US", "currency": "USD", "cabin_category": category,
        "cabin_subcategory": sub or category.upper(),
        "price_pppn": price, "price_per_person": price * nights if price else None,
        "price_total": price * nights * 2 if price else None,
        "availability_status": status, "tier": "weekly-full",
    }


def db(tmp_path, rows):
    path = str(tmp_path / ("p%d.sqlite" % next(_ids)))
    conn = sqlite3.connect(path)
    conn.executescript(DDL)
    for r in rows:
        conn.execute("INSERT INTO observations (%s) VALUES (%s)"
                     % (", ".join(r), ", ".join("?" * len(r))), list(r.values()))
    conn.commit()
    conn.row_factory = sqlite3.Row
    return conn, path


def grid_row(res, line=NCL, region="Caribbean", sail_month="2031-03"):
    return next(r for r in res.rows if r["line"] == line
                and r["region"] == region and r["sail_month"] == sail_month)


def cohort_row(res, scrape, line=NCL, region="Caribbean", cohort="2031-03"):
    return next(r for r in res.rows if r["line"] == line and r["region"] == region
                and r["cohort"] == cohort and r["scrape_date"] == scrape)


# -- the aggregation itself --------------------------------------------------

class TestMinOverMonth:
    def test_min_spans_every_sailing_and_every_scrape_in_the_price_month(
            self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=150.0),
            cell(sailing="B", scrape="2030-05-04", price=140.0),
            cell(sailing="A", scrape="2030-05-18", price=130.0),   # the min
            cell(sailing="B", scrape="2030-05-18", price=145.0),
        ])
        r = grid_row(rj.jefferies_grid(conn, tier="weekly-full"))
        assert r["min_pppn 2030-05"] == 130.0
        assert r["sailings 2030-05"] == 2

    def test_each_price_month_is_its_own_column(self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=200.0),
            cell(sailing="A", scrape="2030-06-01", price=180.0),
            cell(sailing="A", scrape="2030-06-15", price=170.0),
        ])
        r = grid_row(rj.jefferies_grid(conn, tier="weekly-full"))
        assert r["min_pppn 2030-05"] == 200.0
        assert r["min_pppn 2030-06"] == 170.0

    def test_change_is_against_the_first_price_month_observed(self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=200.0),
            cell(sailing="A", scrape="2030-06-01", price=180.0),
        ])
        r = grid_row(rj.jefferies_grid(conn, tier="weekly-full"))
        assert r["first_price_month"] == "2030-05"
        assert r["chg_vs_first 2030-05"] == 0.0
        assert r["chg_vs_first 2030-06"] == -10.0

    def test_a_row_first_seen_later_is_indexed_to_its_own_first_month(
            self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="X", scrape="2030-05-04", price=100.0,
                 sail_date="2031-04-10"),
            cell(sailing="A", scrape="2030-06-01", price=200.0),
            cell(sailing="A", scrape="2030-07-01", price=220.0),
        ])
        r = grid_row(rj.jefferies_grid(conn, tier="weekly-full"))
        assert r["first_price_month"] == "2030-06"
        assert r["min_pppn 2030-05"] is None
        assert r["chg_vs_first 2030-07"] == 10.0

    def test_months_to_sail(self, tmp_path):
        conn, _ = db(tmp_path, [cell(sailing="A", scrape="2030-05-04", price=1.0)])
        r = grid_row(rj.jefferies_grid(conn, tier="weekly-full"))
        assert r["months_to_sail 2030-05"] == 10          # 2031-03 minus 2030-05

    def test_the_grid_names_the_sailing_that_set_the_minimum(self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=150.0, ship="Big"),
            cell(sailing="B", scrape="2030-05-04", price=120.0, ship="Small",
                 sail_date="2031-03-20", nights=4),
        ])
        r = grid_row(rj.jefferies_grid(conn, tier="weekly-full"))
        assert r["min_set_by 2030-05"].startswith("Small 2031-03-20 4n")

    def test_provenance_agrees_with_the_grid(self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=150.0),
            cell(sailing="B", scrape="2030-05-18", price=120.0, ship="Small"),
        ])
        p = rj.min_provenance(conn, tier="weekly-full", line=NCL,
                              region="Caribbean", sail_month="2031-03",
                              price_month_="2030-05")
        assert p["price_pppn"] == 120.0 and p["ship"] == "Small"
        assert p["scrape_date"] == "2030-05-18"


# -- naive vs matched: what the minimum confuses -----------------------------

class TestNaiveVersusMatchedMin:
    def test_a_cheap_sailing_entering_moves_naive_but_not_matched(self, tmp_path):
        """The sell-side minimum reads a newly listed cheap sailing as a cut."""
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=200.0),
            cell(sailing="A", scrape="2030-05-11", price=200.0),
            cell(sailing="NEW", scrape="2030-05-11", price=100.0),
        ])
        r = cohort_row(rj.cohort_index_jefferies(conn, tier="weekly-full"),
                       "2030-05-11")
        assert r["index_naive_min"] == 50.0
        assert r["index_matched_min"] == 100.0
        assert r["mix_effect_min_pp"] == -50.0
        assert r["entered_cells"] == 1

    def test_the_cheapest_sailing_selling_out_moves_naive_but_not_matched(
            self, tmp_path):
        """The case the whole panel exists for: depletion read as a rise."""
        conn, _ = db(tmp_path, [
            cell(sailing="CHEAP", scrape="2030-05-04", price=100.0),
            cell(sailing="DEAR", scrape="2030-05-04", price=200.0),
            cell(sailing="CHEAP", scrape="2030-05-11", price=None,
                 status="sold_out"),
            cell(sailing="DEAR", scrape="2030-05-11", price=200.0),
        ])
        r = cohort_row(rj.cohort_index_jefferies(conn, tier="weekly-full"),
                       "2030-05-11")
        assert r["index_naive_min"] == 200.0        # "prices doubled"
        assert r["index_matched_min"] == 100.0      # no price moved
        assert r["attrition_pct"] == 50.0

    def test_a_genuine_cut_shows_in_both(self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=200.0),
            cell(sailing="A", scrape="2030-05-11", price=180.0),
        ])
        r = cohort_row(rj.cohort_index_jefferies(conn, tier="weekly-full"),
                       "2030-05-11")
        assert r["index_naive_min"] == r["index_matched_min"] == 90.0

    def test_basket_columns_are_minima_and_mean_columns_are_means(self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=100.0),
            cell(sailing="B", scrape="2030-05-04", price=300.0),
        ])
        r = cohort_row(rj.cohort_index_jefferies(conn, tier="weekly-full"),
                       "2030-05-04")
        assert r["basket_base_pppn"] == 100.0
        assert r["naive_min_pppn_now"] == 100.0
        assert r["mean_pppn_now"] == 200.0

    def test_price_month_column(self, tmp_path):
        conn, _ = db(tmp_path, [cell(sailing="A", scrape="2030-05-04", price=1.0)])
        r = cohort_row(rj.cohort_index_jefferies(conn, tier="weekly-full"),
                       "2030-05-04")
        assert r["price_month"] == "2030-05"
        assert r["jefferies_region"] == "Caribbean/Bahamas"


# -- what counts as a comparable balcony -------------------------------------

class TestFilters:
    def base(self, extra):
        return [cell(sailing="A", scrape="2030-05-04", price=200.0)] + extra

    def minimum(self, tmp_path, extra):
        conn, _ = db(tmp_path, self.base(extra))
        return grid_row(rj.jefferies_grid(conn, tier="weekly-full"))["min_pppn 2030-05"]

    def test_minisuite_cannot_set_the_balcony_minimum(self, tmp_path):
        assert self.minimum(tmp_path, [cell(sailing="A", scrape="2030-05-04",
                                            price=50.0, sub="MINISUITE")]) == 200.0

    def test_suites_are_out(self, tmp_path):
        assert self.minimum(tmp_path, [cell(sailing="B", scrape="2030-05-04",
                                            price=50.0, category="suite")]) == 200.0

    def test_packages_are_out(self, tmp_path):
        assert self.minimum(tmp_path, [cell(sailing="B", scrape="2030-05-04",
                                            price=50.0, package=1)]) == 200.0

    def test_unpriced_and_sold_out_are_out(self, tmp_path):
        assert self.minimum(tmp_path, [
            cell(sailing="B", scrape="2030-05-04", price=None, status="sold_out"),
        ]) == 200.0

    def test_limited_cells_count_when_priced(self, tmp_path):
        assert self.minimum(tmp_path, [cell(sailing="B", scrape="2030-05-04",
                                            price=150.0, status="limited")]) == 150.0

    def test_no_length_filter(self, tmp_path):
        """The sell-side takes whatever was listed; so do we."""
        assert self.minimum(tmp_path, [cell(sailing="B", scrape="2030-05-04",
                                            price=90.0, nights=3)]) == 90.0


class TestRegions:
    def test_mapping(self):
        assert rj.jefferies_region("Caribbean") == "Caribbean/Bahamas"
        for r in ("Bermuda", "Alaska", "Southern Europe", "Northern Europe"):
            assert rj.jefferies_region(r) == r
        assert rj.jefferies_region("Transatlantic") is None

    def test_australia_is_declared_not_collected(self):
        assert "Australia/New Zealand" in rj.NOT_COLLECTED

    def test_region_table_lists_uncollected_regions(self, tmp_path):
        conn, _ = db(tmp_path, [cell(sailing="A", scrape="2030-05-04", price=1.0)])
        table = rj.region_table(conn, "weekly-full")
        aus = [r for r in table if r["jefferies_region"] == "Australia/New Zealand"]
        assert aus and aus[0]["note"] == "not collected"

    @pytest.mark.parametrize("sm,pm,n", [("2031-03", "2030-05", 10),
                                         ("2030-05", "2030-05", 0),
                                         ("2030-01", "2029-12", 1)])
    def test_months_between(self, sm, pm, n):
        assert rj.months_between(sm, pm) == n


# -- min vs mean -------------------------------------------------------------

class TestMinVsMean:
    def test_ratio_is_across_sailings_in_a_sail_month(self, tmp_path):
        conn, _ = db(tmp_path, [
            cell(sailing="A", scrape="2030-05-04", price=100.0,
                 sail_date="2031-03-10"),
            cell(sailing="B", scrape="2030-05-04", price=200.0,
                 sail_date="2031-03-12"),
        ])
        res = rj.min_vs_mean_by_lead(conn, tier="weekly-full")
        r = res.rows[0]
        assert r["median_mean_over_min"] == 1.5         # mean 150 / min 100
        assert r["sailings"] == 2

    def test_single_sailing_groups_are_excluded_and_counted(self, tmp_path):
        """A one-sailing group is 1.0 by construction and would drag the
        median toward 1 if it were allowed in."""
        conn, _ = db(tmp_path, [cell(sailing="A", scrape="2030-05-04",
                                     price=100.0)])
        r = rj.min_vs_mean_by_lead(conn, tier="weekly-full").rows[0]
        assert r["median_mean_over_min"] is None
        assert r["single_sailing_groups_excluded"] == 1


# -- the workbook ------------------------------------------------------------

class TestWorkbook:
    def rows(self):
        return [cell(sailing="A", scrape="2030-05-04", price=200.0),
                cell(sailing="A", scrape="2030-05-11", price=190.0),
                cell(sailing="B", line=CCL, scrape="2030-05-11", price=150.0,
                     sub="OB")]

    def test_writes_a_new_workbook_named_for_the_latest_scrape(self, tmp_path):
        _, path = db(tmp_path, self.rows())
        s = rj.write_report(path, out_dir=str(tmp_path / "r"))
        assert os.path.basename(s["path"]) == \
            "2030-05-11__weekly-full__jefferies-basis.xlsx"
        wb = openpyxl.load_workbook(s["path"])
        assert wb.sheetnames == ["Index", "Cohort index (Jefferies basis)",
                                 "Jefferies grid", "Min vs mean by lead"]

    def test_index_documents_the_basis(self, tmp_path):
        _, path = db(tmp_path, self.rows())
        s = rj.write_report(path, out_dir=str(tmp_path / "r"))
        text = "\n".join(str(c.value) for row in
                         openpyxl.load_workbook(s["path"])["Index"].iter_rows()
                         for c in row if c.value is not None)
        for needle in ("Taxes, fees, port charges", "EXCLUDED", "Occupancy",
                       "Cruise-only vs packages", "Australia/New Zealand",
                       "Caribbean/Bahamas"):
            assert needle in text, needle

    def test_an_unchanged_rerun_does_not_rewrite(self, tmp_path):
        _, path = db(tmp_path, self.rows())
        out = str(tmp_path / "r")
        first = rj.write_report(path, out_dir=out)
        mtime = os.path.getmtime(first["path"])
        again = rj.write_report(path, out_dir=out)
        assert not again["written"]
        assert os.path.getmtime(first["path"]) == mtime

    def test_empty_tier_writes_nothing(self, tmp_path):
        _, path = db(tmp_path, [])
        s = rj.write_report(path, out_dir=str(tmp_path / "r"))
        assert s["path"] is None
